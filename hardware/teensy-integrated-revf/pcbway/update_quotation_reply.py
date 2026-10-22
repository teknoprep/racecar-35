#!/usr/bin/env python3
"""Answer PCBWay's BOM quotation IN THEIR OWN WORKBOOK.

PCBWay (2026-10-22, 02:45) corrected the earlier mix-up and sent
`Quotation for T-2SJ3W1113248A-1unit-Racecar-RevF-BOM-pcbway-assembly(2026-10-22).xls`
with: "please carefully check the attachment, we will just follow the final BOM list."

So this does NOT build our own layout (that is what `make_pcbway_quotation_reply.py`
did for the previous round). It updates THEIR sheet in place:

* `Actual Purchase Mfg Part #` (column K) filled on every line with our exact MPN,
* `Customer Reply` (column M) filled on the lines that carry a note or need a
  decision (D24, JP1/JP2, U1, BT1, U21, D23, U8, TRACO U3/U6),
* an extra `Customer Reply` sheet with the build spec, the price/lead-time
  answer and the questions we still need answered,
* saves as `...-REPLIED.xlsx` and `...-REPLIED.xls` next to the original.

It HARD-FAILS if the quotation's 77 lines do not match
`Racecar-RevF-BOM-pcbway-assembly.csv` (MPN / manufacturer / designators / qty),
because "we will just follow the final BOM list" makes that list the contract.
The one expected difference is U1: PCBWay correctly raised its quantity to 2
(two 1x24 sockets in the one 2x24 U1 footprint) - that is reported and accepted.

Usage: python3 pcbway/update_quotation_reply.py [quotation.xls]
       (default: the NEWEST Quotation for T-2SJ3W... RevF .xls, by mtime)
"""
import csv
import datetime
import pathlib
import shutil
import subprocess
import sys
import tempfile

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HERE = pathlib.Path(__file__).resolve().parent
HARDWARE = HERE.parents[1]
REVF_BOM = HERE / 'Racecar-RevF-BOM-pcbway-assembly.csv'

SRC = (pathlib.Path(sys.argv[1]) if len(sys.argv) > 1 else
       # NEWEST by mtime - PCBWay reissues the same project under a new date /
       # quantity, so sorting by filename picks the OLD 5-unit sheet.
       max((p for p in HARDWARE.glob('Quotation for *T-2SJ3W1113248A*RevF*.xls')
            if 'REPLIED' not in p.name),
           key=lambda p: p.stat().st_mtime))
STEM = SRC.stem
OUT_XLSX = HARDWARE / (STEM + '-REPLIED.xlsx')
OUT_XLS = HARDWARE / (STEM + '-REPLIED.xls')

# --- per-line replies -------------------------------------------------------
# keyed by the quotation's Item # (identical to the No. in our BOM CSV)
REPLIES = {
    3: ('Yes - BZX84-C15,215 (Nexperia) is correct, please supply that exact part. '
        'It is the gate clamp on the reverse-voltage FET.'),
    7: ('Must be the SMB (DO-214AA) 5 A / 60 V part: BORN SS56 SMB (LCSC C2687867) or '
        'Diotec SK56. Do NOT use the SMA / DO-214AC SS56 - it does not fit the land pattern.'),
    17: ('Genuine TRACO Power TSR 1-2433 only - no clone or second-source switching '
         'regulator. Both U3 and U6 are this part.'),
    26: ('Confirmed: fit the two 2-pin headers only (TSW-102-07-G-S). Do NOT fit the '
         '2.54 mm shunts - we fit the JP1 shunt after the rail tests and the JP2 shunt '
         'only when this board is an end of the CAN bus.'),
    27: ('Confirmed: quantity 2 is correct. U1 is ONE 2x24 socket footprint and takes two '
         '1x24 SSW-124-01-G-S strips (one per row). Do not source or fit a Teensy 4.1 - '
         'it is customer supplied (bought with pins) and plugs in.'),
    28: ('Correct: 1058 is the Keystone 20 mm coin-cell HOLDER. Fit the holder only - the '
         'CR2032 CELL is customer supplied (primary cell, fitted after cleaning, never charged).'),
    30: ('Correct: order TCAN1042HGVDRQ1. The other part numbers in the description are '
         'fallbacks ONLY and need our written approval before use. The "V" suffix is '
         'mandatory (it is the VIO pin 5); a non-V TCAN1042 has pin 5 = NC and would put '
         '5 V on the 3.3 V Teensy. No clones, no re-marked parts; authorized distribution '
         'only. We inspect the top marking on receipt.'),
    41: ('Accepted - quote the real order-day price. Please tell us the delta BEFORE you '
         'place the order if it moves by more than ~10%.'),
}

SPEC = """RACECAR-35 REV F - CONFIRMED BUILD SPECIFICATION
(2026-10-22, 1 assembled unit)

GERBER: uploaded to the [under review] slot - Racecar-RevF-PCBWAY-GERBERS.zip,
16 files, COMPLETE set (4 copper layers, both masks, both pastes, both
silkscreens, Edge_Cuts, PTH + NPTH drills with maps, .gbrjob). This includes the
updated silkscreen (connector names, pin labels, Teensy orientation, crossed
screen-lead table). Nothing here is a one-layer upload.

STACK / PROCESS
  FR-4, 4 layers, 150.05 x 155.05 mm, 1.6 mm
  1 oz outer copper (inner layers per your standard 4-layer stackup)
  ENIG, green solder mask, white silkscreen
  Impedance control ON - 50 ohm single-ended on the GNSS feed
  Min track 0.20 mm / min spacing 0.15 mm / min hole 0.30 mm
  NO panelisation (connectors and the coin-cell holder sit on the edges)
  The 48 x 6 mm notch in the bottom edge is profile routing (Edge_Cuts),
  not an internal slot.

BOM
  77 lines, and the attached sheet matches our Rev F BOM line for line:
  same MPN, manufacturer, designators and quantities on every line, with ONE
  intended correction - U1 quantity 2 (two 1x24 sockets), which is right.
  BOM and CPL correspond 1:1: 182 designators each. H1-H4 / TP1-TP7 are PCB
  features (NPTH mounting holes, bare-copper test pads): no part, no placement.

  Actual Purchase Mfg Part # = column K, filled on all 77 lines. There is one
  exact part per line - no "equivalent" substitutions. Critical:
    U21  TCAN1042HGVDRQ1 (TI) - V suffix mandatory, authorized distribution,
         NO substitution (see the Customer Reply on that line)
    D23  SS56 - SMB / DO-214AA only
    U1   2x SSW-124-01-G-S sockets, no Teensy
    JP1/JP2  headers only, no shunts
    BT1  1058 holder only, no CR2032 cell
    J9   B2B-XH-A(LF)(SN) connector only, no lead
    U3/U6  genuine TRACO TSR 1-2433

COMMERCIAL
  We accept the 1-unit component + assembly prices shown (169.669 + 29.00).
  Please now add the PCB price (your 5-piece minimum is fine - tell us if that
  changes the unit PCB price) so the All Total is complete, and give us the
  longest lead time on the sheet. If it is easy, also show the assembly total
  at 5 units for reference - we may order more boards than assembled units.
  When the Gerber check is done, please send the proforma so we can pay.

  Reply channel: this sheet's Customer Reply column (M) and this tab.
"""


def convert(src: pathlib.Path, fmt: str, outdir: pathlib.Path) -> pathlib.Path:
    subprocess.run(['libreoffice', '--headless', '--convert-to', fmt,
                    '--outdir', str(outdir), str(src)],
                   check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    return outdir / (src.stem + '.' + fmt)


# --- our BOM ----------------------------------------------------------------
ours = {}
with REVF_BOM.open() as f:
    for row in csv.DictReader(f):
        ours[int(row['No.'])] = dict(des=row['Designator'], qty=int(row['Quantity']),
                                     maker=row['Manufacturer'],
                                     mpn=row['Manufacturer Part Number'])

tmp = pathlib.Path(tempfile.mkdtemp(prefix='pcbway-reply-'))
try:
    xlsx = convert(SRC, 'xlsx', tmp)
    wb = openpyxl.load_workbook(xlsx)
    ws = wb.worksheets[0]

    # header = the row whose A is 'Item #'
    hdr = next(r for r in range(1, 30) if str(ws.cell(r, 1).value).strip() == 'Item #')
    items, r = [], hdr + 1
    while isinstance(ws.cell(r, 1).value, (int, float)):
        items.append(r)
        r += 1
    if len(items) != len(ours) or len(items) != 77:
        raise SystemExit(f'ABORT: {len(items)} item rows in the quotation, {len(ours)} in our BOM')

    # --- verify their list is our BOM, then fill K/M ------------------------
    problems, u1_note = [], False
    for row in items:
        no = int(ws.cell(row, 1).value)
        q = ours.get(no)
        if q is None:
            problems.append(f'item {no}: not in our BOM')
            continue
        des, qty = str(ws.cell(row, 2).value), int(ws.cell(row, 3).value)
        maker, mpn = str(ws.cell(row, 4).value), str(ws.cell(row, 5).value)
        if mpn != q['mpn']:
            problems.append(f'item {no}: MPN {mpn!r} != {q["mpn"]!r}')
        if maker != q['maker']:
            problems.append(f'item {no}: maker {maker!r} != {q["maker"]!r}')
        if des.split() != q['des'].split():
            problems.append(f'item {no}: designators {des!r} != {q["des"]!r}')
        if qty != q['qty']:
            if no == 27 and qty == 2 and q['mpn'] == 'SSW-124-01-G-S':
                u1_note = True          # the intended correction: two sockets
            else:
                problems.append(f'item {no}: qty {qty} != {q["qty"]}')
        ws.cell(row, 11).value = q['mpn']            # Actual Purchase Mfg Part #
        if no in REPLIES:
            c = ws.cell(row, 13)
            c.value = REPLIES[no]
            c.alignment = Alignment(wrap_text=True, vertical='top')
            c.fill = PatternFill('solid', fgColor='FFF2CC')
    if problems:
        raise SystemExit('ABORT: the quotation does not match our BOM:\n  ' +
                         '\n  '.join(problems))

    # --- summary tab ---------------------------------------------------------
    sh = wb.create_sheet('Customer Reply')
    col = 1
    sh.column_dimensions['A'].width = 118
    for i, line in enumerate(SPEC.splitlines(), start=1):
        c = sh.cell(i, col, line)
        c.font = Font(name='Consolas', size=9, bold=(i == 1))
    sh.cell(len(SPEC.splitlines()) + 2, col,
            'Lines answered in column M: ' +
            ', '.join(f'item {n} ({ours[n]["des"]})' for n in sorted(REPLIES)))
    wb.save(OUT_XLSX)
finally:
    shutil.rmtree(tmp, ignore_errors=True)

# --- .xls (the format PCBWay sent) ------------------------------------------
outdir = pathlib.Path(tempfile.mkdtemp(prefix='pcbway-xls-'))
try:
    convert(OUT_XLSX, 'xls', outdir)
    shutil.copy(outdir / (OUT_XLSX.stem + '.xls'), OUT_XLS)
finally:
    shutil.rmtree(outdir, ignore_errors=True)

print(f'wrote {OUT_XLSX.name} and {OUT_XLS.name}')
print(f'  77 lines checked against {REVF_BOM.name}: match'
      + (' (U1 qty 2 = the intended socket correction)' if u1_note else ''))
print('  column K filled on all 77 lines; column M answered on: '
      + ', '.join(str(n) for n in sorted(REPLIES)))
