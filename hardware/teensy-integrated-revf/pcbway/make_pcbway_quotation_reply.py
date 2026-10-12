#!/usr/bin/env python3
"""Build the Rev F answer to PCBWay's quotation T-2SJ3W1113248A.

PCBWay (Kelly Fu, 2026-10-07) asked us to reply to the six notes in the
`*PCBWay Note` column of their quotation, in the `Customer Reply` column, before
they change the quantity. That quotation was built from our **superseded Rev E
BOM**, so this script rewrites the same workbook against the **Rev F** BOM
(`Racecar-RevF-BOM-pcbway-assembly.csv`, the one we actually want quoted):

* the 77 Rev F BOM lines in PCBWay's own column layout,
* the `Customer Reply` column filled in for every line that needs an answer,
* PCBWay's existing prices/lead times carried forward where the part is unchanged
  (reference only - `Component Cost` is deliberately marked INCOMPLETE),
* two extra sheets: the answers to their six questions, and the Rev E -> Rev F delta.

Output: `hardware/Quotation for T-2SJ3W1113248A-5units-Racecar-RevF-BOM-pcbway-assembly(<date>).xlsx`
and the same file converted to `.xls` (the format PCBWay sent us).

Usage: python3 pcbway/make_pcbway_quotation_reply.py [--date YYYY-MM-DD] [--qty 5]
"""
import argparse
import csv
import datetime
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HERE = Path(__file__).resolve().parent            # hardware/teensy-integrated-revf/pcbway
REVF_BOM = HERE / 'Racecar-RevF-BOM-pcbway-assembly.csv'
HARDWARE = HERE.parents[1]                        # hardware/
REV = 'RevF'

ap = argparse.ArgumentParser()
ap.add_argument('--date', default=datetime.date.today().isoformat())
ap.add_argument('--qty', type=int, default=5)
ap.add_argument('--rev-e-quotation', default=None,
                help='PCBWay quotation .xls built from the Rev E BOM (default: newest one)')
args = ap.parse_args()

if args.rev_e_quotation:
    SRC_XLS = Path(args.rev_e_quotation)
else:
    cands = sorted(HARDWARE.glob('Quotation for *RevE-BOM-pcbway-assembly*.xls'))
    if not cands:
        sys.exit('no Rev E PCBWay quotation found in hardware/ - pass --rev-e-quotation')
    SRC_XLS = cands[-1]

OUT_STEM = f'Quotation for T-2SJ3W1113248A-{args.qty}units-Racecar-{REV}-BOM-pcbway-assembly({args.date})'
OUT_XLSX = HARDWARE / (OUT_STEM + '.xlsx')
OUT_XLS = HARDWARE / (OUT_STEM + '.xls')

# ---------------------------------------------------------------- inputs

def read_xls_rows(path):
    """Return {item_no: {col: value}} for A..N of the first sheet of a .xls/.xlsx."""
    tmp = Path(tempfile.mkdtemp(prefix='pcbway-quote-'))
    if path.suffix.lower() == '.xls':
        subprocess.run(['libreoffice', '--headless', '--convert-to', 'xlsx',
                        '--outdir', str(tmp), str(path)],
                       check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
        xlsx = tmp / (path.stem + '.xlsx')
    else:
        xlsx = path
    ws = openpyxl.load_workbook(xlsx, data_only=True).worksheets[0]
    cols = 'ABCDEFGHIJKLMN'
    rows, header = {}, None
    for r in range(1, ws.max_row + 1):
        a = ws[f'A{r}'].value
        if r == 2:
            header = [str(ws[f'{c}{r}'].value or '').strip() for c in cols]
        if not isinstance(a, (int, float)):
            continue
        v = {c: ws[f'{c}{r}'].value for c in cols}
        if any(str(x).strip() for x in v.values() if x is not None):
            rows[int(a)] = v
    shutil.rmtree(tmp, ignore_errors=True)
    return rows, header


rev_e, rev_e_header = read_xls_rows(SRC_XLS)

rev_f = []
with REVF_BOM.open() as f:
    for row in csv.DictReader(f):
        rev_f.append(dict(no=int(row['No.']), qty=int(row['Quantity']), des=row['Designator'],
                          comment=row['Comment'], fp=row['Footprint'],
                          mpn=row['Manufacturer Part Number'], maker=row['Manufacturer']))

# ---------------------------------------------------------------- match Rev E <-> Rev F
# designator overlap, preferring rows whose MPN is also identical
def des_set(s):
    return set(str(s or '').split())

match, claimed = {}, set()
for rf in rev_f:
    best, best_score = None, 0
    for no, re_row in rev_e.items():
        if no in claimed:
            continue
        inter = des_set(rf['des']) & des_set(re_row['B'])
        if not inter:
            continue
        jac = len(inter) / len(des_set(rf['des']) | des_set(re_row['B']))
        score = 1 + (1 if jac >= 0.5 else 0) + (1 if str(re_row['E'] or '') == rf['mpn'] else 0)
        if jac < 0.5:
            score = 0
        if score > best_score:
            best, best_score = no, score
    if best:
        match[rf['no']] = best
        claimed.add(best)

# J13 used to share a line with the deleted J7 header (Rev E item 30)
for rf in rev_f:
    if rf['no'] not in match and 'J13' in des_set(rf['des']):
        match[rf['no']] = 30

# ---------------------------------------------------------------- Customer Reply text
# keyed by a designator that occurs once in the Rev F BOM
REPLIES = {
    'D24': (
        'APPROVED. BZX84-C15,215 is the correct orderable Nexperia MPN for the same device. '
        'Please supply it.'),
    'D23': (
        'CONFIRMED: SMB (DO-214AA) 5 A / 60 V ONLY. You are right that the bare "SS56" is normally '
        'SMA (DO-214AC) - our land pattern is DO-214AA, so the SMA part will NOT fit. Please fit '
        'BORN "SS56" SMB (LCSC C2687867) or Diotec SK56. Do NOT fit the DO-214AC version.'),
    'U3': (
        'Genuine TRACO TSR 1-2433 only - no equivalent or clone. (Both U3 and U6.)'),
    'U1': (
        'CUSTOMER-SUPPLIED TEENSY - do NOT source and do NOT fit a Teensy, and no substitutes are '
        'needed. We buy the Teensy 4.1 WITH PINS and plug it in; please source and fit '
        '2 x SSW-124-01-G-S (1x24, 2.54 mm, THT female socket). NOTE: this line shows Qty 1 because '
        'U1 is a single placement reference in the CPL - it is TWO sockets per board (10 sockets for '
        '5 boards).'),
    'BT1': (
        '"1058" is correct - it is the Keystone SMD 20 mm coin-cell HOLDER. "CR2032" in the '
        'description is the CELL, which is not in the BOM. Please supply the holder and do NOT fit '
        'the CR2032 cell (primary lithium cell - customer fits it by hand after reflow/cleaning).'),
    'J9': (
        'Fit the JST B2B-XH-A(LF)(SN) connector only. The 2-wire lead to the Teensy VBAT/GND pads '
        'is hand-made by the customer - please do not source it.'),
    'U21': (
        'MUST NOT BE SUBSTITUTED. Order TCAN1042HGVDRQ1 from AUTHORIZED DISTRIBUTION ONLY (TI / '
        'Mouser / Digi-Key / LCSC-original) - no clones, no re-marked parts, no substitution without '
        'our written approval. The "V" suffix is MANDATORY: it is the VIO pin (pin 5); a non-V '
        'TCAN1042 (TCAN1042DRQ1 / HD / GD) has pin 5 = NC and would put 5 V on the 3.3 V Teensy. '
        'Approved alternates, in order: TCAN1042VDRQ1, then NXP TJA1051T/3/1J (its pin 8 "S" must '
        'also go to GND). Never an SN65HVD23x, a plain TJA1051T (no /3), a TJA1050 or a '
        '"pin-compatible" clone. Please send us the top marking / supplier before fitting.'),
    'J14': (
        'J14 = CAN: 1 CANH / 2 CANL / 3 GND. The other three (J4 OIL, J11 THROTTLE, J12 BRAKE) are '
        '1 = +5 V / 2 = return / 3 = signal. Same Phoenix 1729021 as before.'),
    'D25': 'New in Rev F - onsemi NUP2105LT1G (dual CAN bus TVS).',
    'R58': (
        'New in Rev F - split CAN termination legs, 60.4 R 1 %. The 120.8 R termination is enabled '
        'by the JP2 shunt only.'),
    'JP1': (
        'FIT THE 2-PIN HEADERS ONLY (TSW-102-07-G-S, JP1 + JP2) - do NOT fit the 2.54 mm shunts. '
        'The customer fits the JP1 shunt after rail tests and the JP2 shunt only when this board is '
        'an end of the CAN bus.'),
    'U8': (
        'NEO-M9N-00B (genuine u-blox). Your price has moved before - please confirm the final price '
        'before production.'),
    'J8': (
        'PART CHANGED: 901-143 -> 901-143-6RFX (Amphenol RF), the brass-body variant of the same '
        'SMA right-angle THT jack - identical footprint, identical PCB hole pattern (1 x 1.5 mm '
        'centre + 4 x 1.7 mm at +/-2.54 mm), same 50 ohm / 500 mating cycles. Your $41.123 each was '
        'about 1.5-2x the distributor price for the stainless version; the -6RFX is about $6.84 at '
        '5 pcs. PLEASE RE-PRICE.'),
    'D5': (
        'PESD5V0F1BL,315 (Nexperia) - this is the replacement already agreed. Please do not revert '
        'to the unavailable PESD3V3U1UL,315.'),
    'F3': '1206L010/60WR as previously agreed (not the /30YR).',
    'D2': 'onsemi SS14 as previously agreed.',
    'U2': (
        'Rev F keeps the on-board TPS54560BDDAR buck. There is no Pololu module and NO socket at U2.')
}

# parts whose Rev E price must NOT be carried over even though the MPN string is the same
# (D23 is the same type number in a different package: SMA then, SMB now)
FORCE_REPRICE = {'D23'}


def reply_for(rf):
    for d in rf['des'].split():
        if d in REPLIES:
            return REPLIES[d]
    return ''


def needs_price(rf):
    return any(d in FORCE_REPRICE for d in rf['des'].split())

# their notes: carry a note to the Rev F line whenever the line is the same part/refs
def carried(rf, col, unchanged_only=True):
    no = match.get(rf['no'])
    if not no:
        return None
    re_row = rev_e[no]
    if unchanged_only and str(re_row['E'] or '') != rf['mpn']:
        return None
    return re_row[col]

# ---------------------------------------------------------------- Rev E -> Rev F delta
DELTA_ADD = [
    ('U21', 'TCAN1042HGVDRQ1', 'Texas Instruments', 'CAN transceiver, VIO = 3.3 V, STB tied to GND'),
    ('D25', 'NUP2105LT1G', 'onsemi', 'Dual CAN bus TVS'),
    ('J14', '1729021', 'Phoenix Contact', 'CAN field terminal 1 CANH / 2 CANL / 3 GND'),
    ('JP2', 'TSW-102-07-G-S', 'Samtec', 'CAN termination jumper, 2-pin header (no shunt fitted)'),
    ('R58 R59', 'RC0805FR-0760R4L', 'Yageo', 'Split termination legs, 60.4 R 1 %'),
    ('C49 C50', 'GRM21BR71H104KA01L', 'Murata', '100 nF decoupling at U21 VCC / VIO'),
    ('C51', 'GRM21BR71H105KA12L', 'Murata', '1 uF bulk on U21 VCC'),
    ('C52', 'CL21B472KBANNNC', 'Samsung', '4.7 nF at the termination midpoint'),
    ('TP6 TP7', '-', '-', 'CAN_TX / CAN_RX test pads - PCB features, no part, no placement'),
]
DELTA_REMOVE = [
    ('J7', 'B4B-XH-A(LF)(SN)', 'JST',
     'External CAN-module header deleted - the transceiver is soldered on the board now'),
    ('C8', 'GRM21BR71H104KA01L', 'Murata', 'Deleted; C49/C50 take its place around U21'),
    ('H1 H2 H3 H4', 'M3 insulated hardware', '-',
     'Never a component: 3.2 mm non-plated mounting holes. Not in the BOM at all since Rev F.'),
    ('TP1 - TP5', 'PCB test pad', '-',
     'Never a component: 2.0 mm plated through-hole pads. Not in the BOM at all since Rev F.'),
]
DELTA_CHANGE = [
    ('U1', 'TEENSY41 (PJRC)', 'SSW-124-01-G-S (Samtec)',
     'Teensy 4.1 is CUSTOMER SUPPLIED. The assembly line is the two 1x24 sockets only.'),
    ('J8', '901-143', '901-143-6RFX',
     'Same Amphenol jack, brass body. Please re-price (see Customer Reply).'),
    ('D23', '"SS56" SMA (DO-214AC)', 'SMB (DO-214AA) 5 A / 60 V',
     'Must be the SMB part: BORN SS56 SMB (LCSC C2687867) or Diotec SK56.'),
    ('JP1', 'JP1 only', 'JP1 + JP2',
     'JP2 (CAN termination) is a second 2-pin header on the same line; headers only, no shunts.'),
    ('C36 (1u)', 'C36', 'C51 C36',
     'C51 is the new U21 bulk capacitor; C36 unchanged.'),
    ('C47 (4.7n)', 'C47', 'C47 C52',
     'C52 is the new CAN termination capacitor; C47 unchanged.'),
    ('100n C_BOOT', '30 pcs (incl. C8)', '31 pcs (C8 out, C49 C50 in)',
     'Same Murata part, quantity adjusted for the CAN transceiver decoupling.'),
    ('D7 - D18 (PMEG)', '7 pcs', '8 pcs (adds D22)',
     'D22 is the new 3.0 V shunt on the car input-voltage monitor.'),
    ('J14 added to J4/J11/J12', '3 x 1729021', '4 x 1729021',
     'J14 is the CAN field terminal - same Phoenix part as the three analog inputs.'),
    ('R16 - R29 (20k 0.1%)', '10 pcs', '11 pcs (adds R49)',
     'R49 is the lower leg of the input-voltage monitor divider; R48 is new too.'),
]

# ---------------------------------------------------------------- workbook
wb = openpyxl.Workbook()
ws = wb.active
ws.title = 'RevF BOM Quotation'

title = (f'PCBWay Bom Quotation, Product No.:T-2SJ3W1113248A  -  Racecar-35 REV F  '
         f'({args.qty} units)  -  THIS REPLACES YOUR REV E QUOTATION OF 2026-10-07')
ws['A1'] = title
ws.merge_cells('A1:N1')
ws['A1'].font = Font(bold=True, size=12, color='FFFFFF')
ws['A1'].fill = PatternFill('solid', fgColor='C00000')
ws['A1'].alignment = Alignment(vertical='center', wrap_text=False)
ws.row_dimensions[1].height = 22

HDR = ['Item #', '*Designator', '*Qty', 'Manufacturer', '*Mfg Part #', 'Description / Value',
       '*Package/Footprint', f'*Unit Price({args.qty} units)', '*Total', '*Delivery Time',
       '*Actual Purchase Mfg Part #', '*PCBWay Note', 'Customer Reply', 'PCBWay Update']
for i, h in enumerate(HDR, 1):
    c = ws.cell(2, i, h)
    c.font = Font(bold=True)
    c.fill = PatternFill('solid', fgColor='FFE699')
    c.alignment = Alignment(vertical='center', wrap_text=True,
                            horizontal='center' if i in (1, 3, 8, 9, 10) else 'left')
ws.row_dimensions[2].height = 30

thin = Side(style='thin', color='B0B0B0')
border = Border(left=thin, right=thin, top=thin, bottom=thin)

priced_total = 0.0
missing_price = []
row = 3
for rf in rev_f:
    price = None if needs_price(rf) else carried(rf, 'H')
    total = round(price * rf['qty'] * args.qty, 3) if price is not None else None
    if total is not None:
        priced_total += total
    else:
        missing_price.append(rf['des'])
    vals = [rf['no'], rf['des'], rf['qty'], rf['maker'], rf['mpn'], rf['comment'], rf['fp'],
            price, total, carried(rf, 'J'), carried(rf, 'K'),
            carried(rf, 'L', unchanged_only=False), reply_for(rf), None]
    for i, v in enumerate(vals, 1):
        c = ws.cell(row, i, v)
        c.border = border
        if i in (1, 3, 8, 9):
            c.alignment = Alignment(horizontal='center', vertical='top')
        elif i == 8 or i == 9:
            c.number_format = '0.000'
        else:
            c.alignment = Alignment(vertical='top', wrap_text=True,
                                    horizontal='left')
    if vals[12]:
        ws.cell(row, 13).font = Font(bold=True)
    row += 1

last_item_row = row - 1
ROW0 = row + 1           # totals block
totals = [
    (f'Component Cost (Rev F, {args.qty} units)',
     round(priced_total, 3),
     f'INCOMPLETE - prices carried forward from your Rev E quotation where the part is '
     f'unchanged. {len(missing_price)} of {len(rev_f)} lines are new or changed and are '
     f'still to be priced: ' + ', '.join(missing_price)),
    ('Assembly Cost', None,
     'Please re-quote for the Rev F BOM/CPL (182 placements, 77 BOM lines).'),
    ('PCB Cost', 174.05,
     'Your Rev E figure - please re-confirm against the Rev F Gerbers (same board, same 150 x '
     '155 mm 4-layer / 1 oz / ENIG / impedance-control settings).'),
    (f'All Total {args.qty} Units', None,
     'Pending - please send the completed Rev F quotation and we will confirm and pay.'),
]
for i, (label, value, note) in enumerate(totals):
    r = ROW0 + i
    ws.cell(r, 8, label).font = Font(bold=True)
    if value is not None:
        c = ws.cell(r, 9, value)
        c.number_format = '0.000'
    else:
        ws.cell(r, 9, '-').alignment = Alignment(horizontal='center')
    ws.cell(r, 10, note).alignment = Alignment(wrap_text=False)

note_r = ROW0 + len(totals) + 1
ws.merge_cells(start_row=note_r, start_column=1, end_row=note_r, end_column=14)
ws.cell(note_r, 1, (
    'HOW TO READ THIS SHEET: the "Customer Reply" column answers every line you flagged, plus the '
    'lines we are changing for Rev F. Our answers to the six questions in your email of 2026-10-07 '
    'are on the "Your questions" sheet, and everything that differs from the Rev E quotation is on '
    'the "Rev E to Rev F" sheet. The BOM/CPL/Gerbers to quote from are attached '
    '(Racecar-RevF-BOM-pcbway-assembly.csv, Racecar-RevF-CPL.csv, Racecar-RevF-PCBWAY-GERBERS.zip). '
    'Every designator in the BOM is in the CPL and vice versa - there is no part to fit for H1-H4 or '
    'TP1-TP7.'))
ws.cell(note_r, 1).font = Font(italic=True, color='C00000')
ws.cell(note_r, 1).alignment = Alignment(wrap_text=True, vertical='top')
ws.row_dimensions[note_r].height = 46

WIDTHS = {'A': 6.5, 'B': 12, 'C': 5.5, 'D': 20, 'E': 22, 'F': 44, 'G': 26,
          'H': 12, 'I': 10, 'J': 14, 'K': 26, 'L': 34, 'M': 62, 'N': 12}
for col, w in WIDTHS.items():
    ws.column_dimensions[col].width = w
ws.freeze_panes = 'A3'
ws.auto_filter.ref = f'A2:N{last_item_row}'

# --- sheet 2: the six questions -------------------------------------------------
q = wb.create_sheet('Your questions')
q['A1'] = 'Answers to the notes in your quotation of 2026-10-07 (quotation T-2SJ3W1113248A)'
q['A1'].font = Font(bold=True, size=12, color='FFFFFF')
q['A1'].fill = PatternFill('solid', fgColor='C00000')
q.merge_cells('A1:E1')
q.row_dimensions[1].height = 22
for i, h in enumerate(['Your item #', 'Ref', 'Your question (verbatim)', 'Our answer',
                       'Rev F item #'], 1):
    c = q.cell(3, i, h)
    c.font = Font(bold=True)
    c.fill = PatternFill('solid', fgColor='FFE699')
    c.alignment = Alignment(wrap_text=True, vertical='center')

rev_f_no = {tuple(r['des'].split()): r['no'] for r in rev_f}
QA = [
    ('3', 'D24 (was item 3)', rev_e[3]['L'], REPLIES['D24'], 3),
    ('7', 'D23 (was item 7)', rev_e[7]['L'], REPLIES['D23'], 7),
    ('27', 'U1 (was item 27)', rev_e[27]['L'], REPLIES['U1'], 27),
    ('28', 'BT1 (was item 28)', rev_e[28]['L'], REPLIES['BT1'], 28),
    ('75', 'H1-H4 (was item 75)', rev_e[75]['L'],
     'NO PART - nothing to source, nothing to place, no assembly cost. H1-H4 are 3.2 mm NON-PLATED '
     'mounting holes in the PCB. The "M3 insulated hardware" note is customer-fitted. These designators '
     'are no longer BOM lines at all in Rev F (they are PCB features, and they stay out of the CPL too), '
     'so your part matching has nothing to stall on.', None),
    ('76', 'TP1-TP5 (was item 76)', rev_e[76]['L'],
     'NO PART - nothing to source, nothing to place, no assembly cost. TP1-TP5 (and the new TP6/TP7) '
     'are 2.0 mm PLATED through-hole pads - bare copper test points. There is no component, no part '
     'number and no brand. Also no longer BOM lines in Rev F.', None),
    ('-', 'J8 (was item 37)',
     'Price flagged at $41.123 each for 901-143.',
     'Part changed to 901-143-6RFX (the brass variant of the same Amphenol jack, identical footprint '
     'and hole pattern). Please re-price - it should be far below $41.123.', 42),
    ('-', 'Delivery-time flags',
     'Your Rev E items 17, 24, 28, 37, 47, 65 and 74 were flagged "7-10 Workdays".',
     'Noted, thank you. Please state the lead time for every line in the Rev F quotation, and tell us '
     'the longest one - we are ordering prototypes, not production.', None),
    ('-', 'U8 (was item 36)', '(no note - price accepted)',
     'NEO-M9N-00B: price accepted, but please confirm the FINAL price before production (it has moved '
     'before).', 41),
    ('-', 'Quantity', 'You asked whether to change the quantity first.',
     f'Please quote {args.qty} units on this Rev F BOM. We will confirm the final quantity once the '
     f'part questions are closed.', None),
]
r = 4
for item, ref, question, answer, rfno in QA:
    q.cell(r, 1, item).alignment = Alignment(horizontal='center', vertical='top')
    q.cell(r, 2, ref).alignment = Alignment(vertical='top', wrap_text=True)
    q.cell(r, 3, question).alignment = Alignment(vertical='top', wrap_text=True)
    q.cell(r, 4, answer).alignment = Alignment(vertical='top', wrap_text=True)
    q.cell(r, 4).font = Font(bold=True)
    q.cell(r, 5, rfno if rfno else '-').alignment = Alignment(horizontal='center', vertical='top')
    for cc in range(1, 6):
        q.cell(r, cc).border = border
    r += 1
for col, w in {'A': 11, 'B': 24, 'C': 52, 'D': 88, 'E': 12}.items():
    q.column_dimensions[col].width = w

# --- sheet 3: the Rev E -> Rev F delta ------------------------------------------
d = wb.create_sheet('Rev E to Rev F')
d['A1'] = ('Racecar-35: what differs between the Rev E quotation you sent and the Rev F BOM we '
           'want built')
d['A1'].font = Font(bold=True, size=12, color='FFFFFF')
d['A1'].fill = PatternFill('solid', fgColor='C00000')
d.merge_cells('A1:D1')
d.row_dimensions[1].height = 22
d.cell(3, 1, ('The board is the same size, the same 4 layers and the same process. Rev E is '
              'SUPERSEDED: Rev F is the revision that fixes CAN (its external CAN-module header J7 '
              'is deleted and a soldered TCAN1042HGV-Q1 replaces the module).')).font = Font(italic=True)

r = 5
for heading, rows, cols in (
        ('ADDED in Rev F', DELTA_ADD, ['Ref', 'MPN', 'Manufacturer', 'What it is']),
        ('REMOVED in Rev F', DELTA_REMOVE, ['Ref', 'Rev E part', 'Manufacturer', 'Why']),
        ('CHANGED in Rev F', DELTA_CHANGE, ['Ref', 'Rev E', 'Rev F', 'Why / note'])):
    c = d.cell(r, 1, heading)
    c.font = Font(bold=True, size=11)
    c.fill = PatternFill('solid', fgColor='FFE699')
    r += 1
    for i, h in enumerate(cols, 1):
        cc = d.cell(r, i, h)
        cc.font = Font(bold=True, italic=True)
    r += 1
    for row_vals in rows:
        for i, v in enumerate(row_vals, 1):
            cc = d.cell(r, i, v)
            cc.alignment = Alignment(vertical='top', wrap_text=(i == len(row_vals)))
            cc.border = border
        r += 1
    r += 1
r += 1
c = d.cell(r, 1, 'LINE NUMBERS: your Rev E quotation item -> our Rev F quotation item')
c.font = Font(bold=True, size=11)
c.fill = PatternFill('solid', fgColor='FFE699')
r += 1
d.cell(r, 1, 'Only rows that moved are shown; every other line kept its number.').font = \
    Font(italic=True)
r += 1
line_map = [(no, rf['no'], rf['des']) for rf in rev_f
            for no in (match.get(rf['no']),) if no and no != rf['no']]
line_map.sort(key=lambda t: t[0])
for i, h in enumerate(['Your Rev E item #', 'Ref', 'Our Rev F item #'], 1):
    d.cell(r, i, h).font = Font(bold=True, italic=True)
r += 1
for no, rfno, des in line_map:
    d.cell(r, 1, no).alignment = Alignment(horizontal='center')
    d.cell(r, 2, des).alignment = Alignment(wrap_text=True)
    d.cell(r, 3, rfno).alignment = Alignment(horizontal='center')
    for cc in range(1, 4):
        d.cell(r, cc).border = border
    r += 1

for col, w in {'A': 22, 'B': 30, 'C': 26, 'D': 78}.items():
    d.column_dimensions[col].width = w

wb.save(OUT_XLSX)
subprocess.run(['libreoffice', '--headless', '--convert-to', 'xls', '--outdir',
                str(HARDWARE), str(OUT_XLSX)],
               check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
print(f'wrote {OUT_XLSX.name} and {OUT_XLS.name}')
print(f'  {len(rev_f)} Rev F BOM lines, {sum(1 for rr in rev_f if reply_for(rr))} with a Customer Reply')
print(f'  prices carried for {len(rev_f) - len(missing_price)} lines; '
      f'{len(missing_price)} to be re-priced: ' + ', '.join(missing_price))
